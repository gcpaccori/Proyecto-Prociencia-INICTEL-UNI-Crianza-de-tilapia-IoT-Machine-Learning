"""Extract published observations; never interpolate biological labels."""
import csv
import argparse
import hashlib
import json
import math
from datetime import date, timedelta
from pathlib import Path
from docx import Document

parser = argparse.ArgumentParser()
parser.add_argument('--kamble-docx', type=Path, required=True)
parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parent/'data')
args=parser.parse_args()
OUT = args.output
OUT.mkdir(parents=True, exist_ok=True)
source = args.kamble_docx
doc = Document(source)
groups = ['Control', 'GLE-5', 'GLE-10', 'SGLE-5', 'SGLE-10', 'MxLE-5', 'MxLE-10']

def save(name, rows):
    with (OUT / name).open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator='\n')
        w.writeheader(); w.writerows(rows)

fish = []
for group, table in zip(groups, doc.tables[7:14]):
    for row_index, row in enumerate(table.rows[2:], 3):
        cells = [c.text.strip() for c in row.cells]
        if cells[0] not in ('R1', 'R2', 'R3') or not cells[1].isdigit():
            continue
        for endpoint, start, day in [('initial', 2, 0), ('final', 5, 60)]:
            fish.append(dict(study_id='zenodo11174309', treatment=group,
                tank_id=f'{group}_{cells[0]}', endpoint=endpoint, day=day,
                sample_row=cells[1], length_cm=float(cells[start]),
                weight_g=float(cells[start+1]), source_table=groups.index(group)+8,
                source_row=row_index, individually_tagged='not_documented',
                source_doi='10.5281/zenodo.11174309'))
assert len(fish) == 1050
save('kamble_2024_endpoint_samples.csv', fish)

env = []
for table_index, variable, unit in [(0,'temperature_c','degC'),(1,'ph','pH'),(2,'do_mg_l','mg/L'),
        (3,'ammonia_n_mg_l','mg N/L'),(4,'nitrite_n_mg_l','mg N/L'),(5,'nitrate_n_mg_l','mg N/L')]:
    week = None
    for row_index, row in enumerate(doc.tables[table_index].rows, 1):
        cells = [c.text.strip() for c in row.cells]
        if cells[0].isdigit(): week = int(cells[0])
        if len(cells)<2 or cells[1] not in ('R1','R2','R3'): continue
        for gi, group in enumerate(groups):
            for period, col in ([('morning',2+gi*2),('evening',3+gi*2)] if table_index<3 else [('weekly',2+gi)]):
                env.append(dict(study_id='zenodo11174309', treatment=group,
                    tank_id=f'{group}_{cells[1]}', week=week, period=period,
                    variable=variable, value=float(cells[col]), unit=unit,
                    source_table=table_index+1, source_row=row_index))
assert len(env)==1701
save('kamble_2024_weekly_environment.csv', env)

intervals=[]
for group in groups:
    for r in ('R1','R2','R3'):
        tank=f'{group}_{r}'
        a=[x['weight_g'] for x in fish if x['tank_id']==tank and x['endpoint']=='initial']
        b=[x['weight_g'] for x in fish if x['tank_id']==tank and x['endpoint']=='final']
        row=dict(study_id='zenodo11174309', treatment=group, tank_id=tank, duration_days=60,
            initial_weight_g=sum(a)/len(a), final_weight_g=sum(b)/len(b),
            n_initial=len(a), n_final=len(b),
            endpoint_type='tank_sample_mean_not_tagged_individual_growth')
        for v in ('temperature_c','ph','do_mg_l'):
            values=[x['value'] for x in env if x['tank_id']==tank and x['variable']==v]
            row['weekly_summary_mean_'+v]=sum(values)/len(values)
        row['environment_resolution']='weekly_morning_evening_means_not_hourly_exposure'
        intervals.append(row)
assert len(intervals)==21
save('kamble_2024_tank_intervals.csv',intervals)

qc=[]
for gi,group in enumerate(groups):
    initial=[x['weight_g'] for x in fish if x['treatment']==group and x['endpoint']=='initial']
    final=[x['weight_g'] for x in fish if x['treatment']==group and x['endpoint']=='final']
    published=float(next(r.cells[2+gi].text for r in doc.tables[6].rows if r.cells[1].text.strip()=='Mean'))
    log_rate=100*(sum(map(math.log,final))/len(final)-sum(map(math.log,initial))/len(initial))/60
    qc.append(dict(treatment=group,published_mean_sgr_pct_day=published,
        diagnostic_sgr_from_marginal_log_weights_60d=log_rate,
        interpretation='rate_discrepancy_requires_resolution_not_paired_individual_growth',
        source_doi='10.5281/zenodo.11174309',published_rate_source_table=7))
save('kamble_2024_rate_consistency_QC.csv',qc)

opiyo=[]
for diet,wi,wf,sgr,dwg in [('LM0',30.46,100.51,1.10,.66),('LM10',30.50,94.07,1.03,.59),('LM15',30.45,121.05,1.32,.89),('COMM',30.48,148.64,1.51,1.13)]:
    opiyo.append(dict(study_id='opiyo2024',treatment=diet,initial_weight_g=wi,final_weight_g=wf,
        duration_days_reported=84,reported_sgr_pct_day=sgr,recomputed_sgr_for_84_days=100*math.log(wf/wi)/84,
        reported_dwg_g_day=dwg,recomputed_dwg_for_84_days=(wf-wi)/84,
        source_doi='10.5281/zenodo.14555064',source_pdf_page=4,source_table=2,
        status='QUARANTINED_DURATION_OR_RATE_INCONSISTENCY',training_allowed=False))
save('opiyo_2024_QUARANTINED_summary.csv',opiyo)

weights=[3.8380,9.2350,16.4613,22.5657,29.0560,38.6327,52.6893,64.1557,79.4193,85.1807,
         101.45,112.76,125.34,129.54,138.7,142.28,147.18,149.69,157.2,159.3,164.23,172.4]
lengths=[6,7.5,9,10,11,12,13.5,14,16,17,18,19,19.5,20,20,21,21,21,21,22,22,24]
garcia=[dict(study_id='garcia2012_piura',cohort_id='piura_2011',date=(date(2011,5,1)+timedelta(days=i*7)).isoformat(),
    days_since_first_sample=i*7,mean_weight_g=w,mean_length_cm=l,n_reported_table=300,
    sample_n_conflict='Spanish_abstract_30_methods_and_table_300',source_table=1,source_pdf_page=4,
    source_doi='10.21503/CienciayDesarrollo.2012.v15i2.05',
    age_at_first_sample_days='unknown_do_not_use_table2_years') for i,(w,l) in enumerate(zip(weights,lengths))]
save('garcia_2012_observed_weekly_means.csv',garcia)
manifest={'observed_endpoint_rows':len(fish),'experimental_tanks':21,'treatments':7,
    'external_studies_normalized':3,'studies_with_growth_rate_QC_issues':2,
    'observed_weekly_cohort_means':22,
    'weekly_environment_summary_cells':len(env),
    'source_docx_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
    'warnings':['No individual identity across endpoints assumed.',
      'Weekly environmental summaries cannot identify hours below an oxygen threshold.',
      'Do not pool different ages into a universal fish weight histogram.',
      'No corrected local biometric weight was accessed or used.',
      'Kamble published SGR means do not reconcile with endpoint log-weight distributions over 60 days; benchmark uses reported endpoints and duration provisionally.',
      'Opiyo interval rates/duration inconsistent: quarantined, excluded from training.'],
    'files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.glob('*.csv')}}
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2))
print(json.dumps(manifest,indent=2))
